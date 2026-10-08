from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func, or_
from pydantic import BaseModel, EmailStr, Field
from pathlib import Path
from datetime import datetime, date, time, timedelta
from uuid import uuid4
import json, re, io, os
import textwrap
from .database import Base, engine, get_db, migrate_schema, User, Job, Application, Resume, ATSResult
from .auth import hash_password, verify_password, token_for, current_user, require_admin
from .services.scoring import analyze

Base.metadata.create_all(engine)
migrate_schema()
app=FastAPI(title='HireX ATS API', version='1.0.0')
app.add_middleware(CORSMiddleware, allow_origins=['http://localhost:5173'], allow_credentials=True, allow_methods=['*'], allow_headers=['*'])
UPLOADS=Path(__file__).resolve().parents[1]/'uploads'; UPLOADS.mkdir(exist_ok=True)
class Registration(BaseModel): name:str=Field(min_length=2,max_length=120); email:EmailStr; password:str=Field(min_length=8,max_length=128); department:str=Field(min_length=1,max_length=120); register_number:str=Field(min_length=1,max_length=80)
class JobInput(BaseModel): title:str; description:str; minimum_score:int=Field(ge=0,le=100); required_skills:list[str]; optional_skills:list[str]=[]; active:bool=True
class Login(BaseModel): email:EmailStr; password:str
class ProfileInput(BaseModel): name:str=Field(min_length=2,max_length=120); department:str|None=Field(default=None,max_length=120); register_number:str|None=Field(default=None,max_length=80)
class PasswordInput(BaseModel): current_password:str; new_password:str=Field(min_length=8,max_length=128); confirm_password:str

def user_json(u): return {'id':u.id,'name':u.name,'email':u.email,'role':u.role,'department':u.department,'register_number':u.register_number,'photo_url':f'/api/profile/photo?v={Path(u.profile_photo).name}' if u.profile_photo else None}
def job_json(j): return {'id':j.id,'title':j.title,'description':j.description,'minimum_score':j.minimum_score,'required_skills':j.required_skills.splitlines(),'optional_skills':j.optional_skills.splitlines(),'active':j.active,'created_at':j.created_at.isoformat()}
def app_json(a): return {'id':a.id,'student_id':a.student_id,'student_name':a.student.name,'student_email':a.student.email,'department':a.student.department,'register_number':a.student.register_number,'photo_url':'/api/profile/photo' if a.student.profile_photo else None,'job_id':a.job_id,'job_title':a.job.title,'status':a.status,'current_stage':a.current_stage,'created_at':a.created_at.isoformat(),'score':a.result.overall_score if a.result else None,'qualification_status':a.result.qualification_status if a.result else 'Pending','has_resume':bool(a.resume),'assessment_at':a.result.created_at.isoformat() if a.result else None}
@app.get('/api/health')
def health(): return {'status':'ok'}
@app.post('/api/auth/register')
def register(data:Registration,db:Session=Depends(get_db)):
    if db.query(User).filter_by(email=data.email.lower()).first(): raise HTTPException(409,'An account with this email already exists.')
    name=data.name.strip(); department=data.department.strip()
    register_number=data.register_number.strip()
    if len(name)<2: raise HTTPException(422,'Full name is required.')
    if not department: raise HTTPException(422,'Department is required.')
    if not register_number: raise HTTPException(422,'Register number is required.')
    if db.query(User).filter(User.role=='student',func.lower(User.register_number)==register_number.lower()).first(): raise HTTPException(409,'That register number is already in use.')
    u=User(name=name,email=data.email.lower(),password_hash=hash_password(data.password),role='student',department=department,register_number=register_number); db.add(u); db.commit(); db.refresh(u)
    return {'access_token':token_for(u),'token_type':'bearer','user':user_json(u)}
@app.post('/api/auth/login')
def login(data:Login,db:Session=Depends(get_db)):
    u=db.query(User).filter_by(email=data.email.lower()).first()
    if not u or not verify_password(data.password,u.password_hash): raise HTTPException(401,'Incorrect email or password.')
    return {'access_token':token_for(u),'token_type':'bearer','user':user_json(u)}
@app.get('/api/auth/me')
def me(u:User=Depends(current_user)): return user_json(u)
@app.get('/api/profile')
def profile(u:User=Depends(current_user)): return user_json(u)
@app.put('/api/profile')
def update_profile(data:ProfileInput,db:Session=Depends(get_db),u:User=Depends(current_user)):
    name=data.name.strip(); department=(data.department or '').strip() or None; register_number=(data.register_number or '').strip() or None
    if not name: raise HTTPException(422,'Full name is required.')
    if u.role=='student' and not department: raise HTTPException(422,'Department is required for student accounts.')
    if u.role=='student' and not register_number: raise HTTPException(422,'Register number is required for student accounts.')
    if register_number and db.query(User).filter(User.id!=u.id,User.role=='student',func.lower(User.register_number)==register_number.lower()).first(): raise HTTPException(409,'That register number is already in use.')
    u.name=name; u.department=department; u.register_number=register_number; db.commit(); db.refresh(u); return user_json(u)
@app.post('/api/profile/password')
def change_password(data:PasswordInput,db:Session=Depends(get_db),u:User=Depends(current_user)):
    if not verify_password(data.current_password,u.password_hash): raise HTTPException(400,'Current password is incorrect.')
    if data.new_password!=data.confirm_password: raise HTTPException(422,'New password and confirmation do not match.')
    if verify_password(data.new_password,u.password_hash): raise HTTPException(422,'Choose a new password different from your current password.')
    u.password_hash=hash_password(data.new_password); db.commit(); return {'message':'Password changed successfully.'}
PROFILE_DIR=UPLOADS/'profile'; PROFILE_DIR.mkdir(exist_ok=True)
PROFILE_TYPES={'.jpg':('image/jpeg',b'\xff\xd8\xff'),'.jpeg':('image/jpeg',b'\xff\xd8\xff'),'.png':('image/png',b'\x89PNG\r\n\x1a\n')}
@app.post('/api/profile/photo')
def upload_profile_photo(file:UploadFile=File(...),db:Session=Depends(get_db),u:User=Depends(current_user)):
    ext=Path(file.filename or '').suffix.lower(); spec=PROFILE_TYPES.get(ext)
    if not spec or file.content_type!=spec[0]: raise HTTPException(400,'Choose a JPG, JPEG, or PNG image.')
    data=file.file.read(3*1024*1024+1)
    if len(data)>3*1024*1024: raise HTTPException(413,'Profile images must be 3 MB or smaller.')
    if not data.startswith(spec[1]): raise HTTPException(400,'The image content does not match its file type.')
    filename=f'{u.id}_{uuid4().hex}{ext}'; destination=PROFILE_DIR/filename; destination.write_bytes(data)
    if u.profile_photo:
        old=Path(u.profile_photo)
        if old.parent.resolve()==PROFILE_DIR.resolve(): old.unlink(missing_ok=True)
    u.profile_photo=str(destination); db.commit(); return user_json(u)
@app.get('/api/profile/photo')
def get_profile_photo(u:User=Depends(current_user)):
    if not u.profile_photo or not Path(u.profile_photo).is_file(): raise HTTPException(404,'No profile photo is available.')
    path=Path(u.profile_photo); ext=path.suffix.lower()
    if path.parent.resolve()!=PROFILE_DIR.resolve() or ext not in PROFILE_TYPES: raise HTTPException(404,'No profile photo is available.')
    return FileResponse(path,media_type=PROFILE_TYPES[ext][0],headers={'Cache-Control':'private, no-store'})
@app.delete('/api/profile/photo')
def delete_profile_photo(db:Session=Depends(get_db),u:User=Depends(current_user)):
    if u.profile_photo:
        path=Path(u.profile_photo)
        if path.parent.resolve()==PROFILE_DIR.resolve(): path.unlink(missing_ok=True)
        u.profile_photo=None; db.commit()
    return user_json(u)
@app.get('/api/jobs')
def jobs(db:Session=Depends(get_db),u:User=Depends(current_user)):
    return [job_json(j) for j in db.query(Job).filter_by(active=True).order_by(Job.created_at.desc()).all()]
@app.post('/api/jobs')
def create_job(data:JobInput,db:Session=Depends(get_db),u:User=Depends(require_admin)):
    if not data.description.strip(): raise HTTPException(400,'Job description is required.')
    if not data.required_skills: raise HTTPException(400,'At least one required skill is required.')
    j=Job(title=data.title,description=data.description,minimum_score=data.minimum_score,required_skills='\n'.join(data.required_skills),optional_skills='\n'.join(data.optional_skills),active=data.active,created_by=u.id); db.add(j); db.commit(); db.refresh(j); return job_json(j)
@app.get('/api/jobs/{jid}')
def get_job(jid:int,db:Session=Depends(get_db),u:User=Depends(current_user)):
    j=db.get(Job,jid)
    if not j or (u.role!='admin' and not j.active): raise HTTPException(404,'Job not found')
    return job_json(j)
@app.put('/api/jobs/{jid}')
def update_job(jid:int,data:JobInput,db:Session=Depends(get_db),u:User=Depends(require_admin)):
    j=db.get(Job,jid)
    if not j: raise HTTPException(404,'Job not found')
    if not data.title.strip(): raise HTTPException(422,'Job title is required.')
    if not data.description.strip(): raise HTTPException(422,'Job description is required.')
    required=[skill.strip() for skill in data.required_skills if skill.strip()]
    if not required: raise HTTPException(422,'At least one required skill is required.')
    values=data.model_dump()
    values['required_skills']=required
    values['optional_skills']=[skill.strip() for skill in data.optional_skills if skill.strip()]
    for k,v in values.items(): setattr(j,k,'\n'.join(v) if isinstance(v,list) else v)
    db.commit(); db.refresh(j); return job_json(j)
@app.delete('/api/jobs/{jid}')
def delete_job(jid:int,db:Session=Depends(get_db),u:User=Depends(require_admin)):
    j=db.get(Job,jid)
    if not j: raise HTTPException(404,'Job not found')
    j.active=False; db.commit(); return {'ok':True}
@app.post('/api/applications')
def apply(job_id:int,db:Session=Depends(get_db),u:User=Depends(current_user)):
    if u.role!='student': raise HTTPException(403,'Candidate access required')
    if not db.get(Job,job_id) or not db.get(Job,job_id).active: raise HTTPException(404,'Job not found')
    if db.query(Application).filter_by(student_id=u.id,job_id=job_id).first(): raise HTTPException(409,'You have already applied for this job.')
    a=Application(student_id=u.id,job_id=job_id); db.add(a); db.commit(); db.refresh(a); return {'id':a.id,'message':'Application submitted. Upload a resume to begin screening.'}
def owned(a,u):
    if not a: raise HTTPException(404,'Application not found')
    if u.role!='admin' and a.student_id!=u.id: raise HTTPException(403,'You can only access your own application.')
@app.get('/api/applications/my')
def my_apps(db:Session=Depends(get_db),u:User=Depends(current_user)):
    q=db.query(Application)
    if u.role!='admin': q=q.filter_by(student_id=u.id)
    return [app_json(a) for a in q.order_by(Application.created_at.desc()).all()]
@app.get('/api/applications/{aid}')
def application(aid:int,db:Session=Depends(get_db),u:User=Depends(current_user)):
    a=db.get(Application,aid); owned(a,u); return app_json(a)
@app.post('/api/resumes/upload')
def upload(application_id:int=Form(...),file:UploadFile=File(...),db:Session=Depends(get_db),u:User=Depends(current_user)):
    a=db.get(Application,application_id); owned(a,u)
    if a.student_id!=u.id and u.role!='admin': raise HTTPException(403,'You can only upload to your application.')
    ext=Path(file.filename or '').suffix.lower()
    if ext not in ['.pdf','.docx'] or file.content_type not in ['application/pdf','application/vnd.openxmlformats-officedocument.wordprocessingml.document','application/octet-stream']: raise HTTPException(400,'Please upload a PDF or DOCX resume.')
    data=file.file.read(8*1024*1024+1)
    if len(data)>8*1024*1024: raise HTTPException(413,'Resume must be smaller than 8 MB.')
    if not data: raise HTTPException(400,'Unable to read this resume. Please upload a valid PDF or DOCX file.')
    safe=f'{u.id}_{application_id}_{datetime.utcnow().timestamp():.0f}{ext}'; path=UPLOADS/safe; path.write_bytes(data)
    try:
        if ext=='.pdf':
            from pypdf import PdfReader
            from io import BytesIO
            text='\n'.join(p.extract_text() or '' for p in PdfReader(BytesIO(data)).pages)
        else:
            from docx import Document
            from io import BytesIO
            doc=Document(BytesIO(data)); text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' '.join(c.text for row in t.rows for c in row.cells) for t in doc.tables)
        if len(text.strip())<30: raise ValueError('empty')
    except Exception:
        path.unlink(missing_ok=True); raise HTTPException(400,'Unable to read this resume. Please upload a valid PDF or DOCX file.')
    old=db.query(Resume).filter_by(application_id=application_id).first()
    if old: Path(old.file_path).unlink(missing_ok=True); db.delete(old); db.flush()
    res=Resume(student_id=a.student_id,application_id=application_id,filename=Path(file.filename).name,file_path=str(path),extracted_text=text); db.add(res); a.status='Resume uploaded'; db.commit(); return {'filename':res.filename,'characters':len(text),'status':'Resume processed'}
@app.post('/api/ats/analyze/{aid}')
def screen(aid:int,db:Session=Depends(get_db),u:User=Depends(current_user)):
    a=db.get(Application,aid); owned(a,u)
    if not a.resume: raise HTTPException(400,'Upload a resume before starting analysis.')
    j=db.get(Job,a.job_id); total,detail=analyze(a.resume.extracted_text,j)
    result=db.query(ATSResult).filter_by(application_id=aid).first()
    if result: result.overall_score=total; result.details=json.dumps(detail); result.qualification_status=detail['qualification_status']; result.created_at=datetime.utcnow()
    else: result=ATSResult(application_id=aid,overall_score=total,details=json.dumps(detail),qualification_status=detail['qualification_status']); db.add(result)
    a.status=detail['qualification_status'].title().replace(' ',' '); a.current_stage='Aptitude Assessment' if total>=j.minimum_score else 'ATS Screening'; db.commit(); return {'application_id':aid,'score':total,**detail}
@app.get('/api/ats/{aid}')
def get_result(aid:int,db:Session=Depends(get_db),u:User=Depends(current_user)):
    a=db.get(Application,aid); owned(a,u)
    if not a.result: raise HTTPException(404,'Analysis is not available yet.')
    return {'application_id':aid,'score':a.result.overall_score,**json.loads(a.result.details)}
def admin_application_query(db,department=None,register_number=None,search=None,job_id=None,status=None,min_score=None,max_score=None,date_from=None,date_to=None):
    q=db.query(Application).join(User,Application.student_id==User.id).join(Job,Application.job_id==Job.id).outerjoin(ATSResult,ATSResult.application_id==Application.id)
    if department: q=q.filter(func.lower(User.department)==department.lower())
    if register_number: q=q.filter(func.lower(User.register_number).contains(register_number.lower()))
    if search:
        term=f'%{search.strip().lower()}%'
        q=q.filter(or_(func.lower(User.name).like(term),func.lower(User.email).like(term),func.lower(func.coalesce(User.register_number,'')).like(term),func.lower(func.coalesce(User.department,'')).like(term),func.lower(Job.title).like(term)))
    if job_id is not None: q=q.filter(Application.job_id==job_id)
    if status:
        if status.lower()=='pending': q=q.filter(ATSResult.id.is_(None))
        else: q=q.filter(func.lower(ATSResult.qualification_status)==status.lower())
    if min_score is not None: q=q.filter(ATSResult.overall_score>=min_score)
    if max_score is not None: q=q.filter(ATSResult.overall_score<=max_score)
    event_date=func.coalesce(ATSResult.created_at,Application.created_at)
    if date_from: q=q.filter(event_date>=datetime.combine(date_from,time.min))
    if date_to: q=q.filter(event_date<datetime.combine(date_to+timedelta(days=1),time.min))
    return q

@app.get('/api/admin/candidates')
def candidates(db:Session=Depends(get_db),u:User=Depends(require_admin),department:str|None=None,register_number:str|None=None,search:str|None=None,job_id:int|None=None,status:str|None=None,min_score:float|None=Query(default=None,ge=0,le=100),max_score:float|None=Query(default=None,ge=0,le=100),date_from:date|None=None,date_to:date|None=None,limit:int=Query(default=25,ge=1,le=100),offset:int=Query(default=0,ge=0)):
    q=db.query(User,func.max(Application.id).label('latest_aid')).outerjoin(Application,Application.student_id==User.id).outerjoin(Job,Job.id==Application.job_id).outerjoin(ATSResult,ATSResult.application_id==Application.id).filter(User.role=='student')
    if department: q=q.filter(func.lower(User.department)==department.lower())
    if register_number: q=q.filter(func.lower(User.register_number).contains(register_number.lower()))
    if search:
        term=f'%{search.strip().lower()}%'; q=q.filter(or_(func.lower(User.name).like(term),func.lower(User.email).like(term),func.lower(func.coalesce(User.register_number,'')).like(term),func.lower(func.coalesce(User.department,'')).like(term),func.lower(func.coalesce(Job.title,'')).like(term)))
    if job_id is not None: q=q.filter(Application.job_id==job_id)
    if status:
        if status.lower()=='pending': q=q.filter(ATSResult.id.is_(None))
        else: q=q.filter(func.lower(ATSResult.qualification_status)==status.lower())
    if min_score is not None: q=q.filter(ATSResult.overall_score>=min_score)
    if max_score is not None: q=q.filter(ATSResult.overall_score<=max_score)
    event_date=func.coalesce(ATSResult.created_at,Application.created_at)
    if date_from: q=q.filter(event_date>=datetime.combine(date_from,time.min))
    if date_to: q=q.filter(event_date<datetime.combine(date_to+timedelta(days=1),time.min))
    total=q.with_entities(User.id).group_by(User.id).count()
    rows=q.group_by(User.id).order_by(func.max(Application.created_at).desc().nullslast(),User.name).offset(offset).limit(limit).all()
    result=[]
    for person,latest_aid in rows:
        selected_aid=latest_aid
        if not any([job_id is not None,status,min_score is not None,max_score is not None,date_from,date_to]):
            latest_screening=db.query(Application.id).join(ATSResult,ATSResult.application_id==Application.id).filter(Application.student_id==person.id).order_by(ATSResult.created_at.desc()).first()
            if latest_screening: selected_aid=latest_screening[0]
        app_row=db.get(Application,selected_aid) if selected_aid else None
        if app_row:
            item=app_json(app_row)
        else:
            item={'student_id':person.id,'student_name':person.name,'student_email':person.email,'department':person.department,'register_number':person.register_number,'job_id':None,'job_title':'No application yet','status':'Registered','current_stage':'Not started','created_at':person.created_at.isoformat(),'score':None,'qualification_status':'Pending','has_resume':False,'assessment_at':None}
        item.update({'photo_url':f'/api/admin/candidates/{person.id}/photo?v={Path(person.profile_photo).name}' if person.profile_photo else None,'applications_count':db.query(Application).filter_by(student_id=person.id).count()}); result.append(item)
    return {'items':result,'total':total,'limit':limit,'offset':offset}

@app.get('/api/admin/candidates/{student_id}')
def candidate_detail(student_id:int,db:Session=Depends(get_db),u:User=Depends(require_admin)):
    person=db.get(User,student_id)
    if not person or person.role!='student': raise HTTPException(404,'Candidate not found.')
    applications=db.query(Application).outerjoin(ATSResult,ATSResult.application_id==Application.id).filter(Application.student_id==student_id).order_by(ATSResult.created_at.desc().nullslast(),Application.created_at.desc()).all()
    results=[]
    for app_row in applications:
        item=app_json(app_row)
        if app_row.result:
            item['analysis']={'application_id':app_row.id,'score':app_row.result.overall_score,'qualification_status':app_row.result.qualification_status,**json.loads(app_row.result.details)}
        results.append(item)
    profile_data=user_json(person); profile_data['photo_url']=f'/api/admin/candidates/{person.id}/photo?v={Path(person.profile_photo).name}' if person.profile_photo else None
    return {'profile':profile_data,'applications':results}
@app.get('/api/admin/candidates/{student_id}/photo')
def admin_candidate_photo(student_id:int,db:Session=Depends(get_db),u:User=Depends(require_admin)):
    person=db.get(User,student_id)
    if not person or person.role!='student': raise HTTPException(404,'Candidate not found.')
    if not person.profile_photo or not Path(person.profile_photo).is_file(): raise HTTPException(404,'No profile photo is available.')
    path=Path(person.profile_photo); ext=path.suffix.lower()
    if path.parent.resolve()!=PROFILE_DIR.resolve() or ext not in PROFILE_TYPES: raise HTTPException(404,'No profile photo is available.')
    return FileResponse(path,media_type=PROFILE_TYPES[ext][0],headers={'Cache-Control':'private, no-store'})

@app.get('/api/admin/analytics')
def analytics(db:Session=Depends(get_db),u:User=Depends(require_admin),department:str|None=None,register_number:str|None=None,search:str|None=None,job_id:int|None=None,status:str|None=None,min_score:float|None=Query(default=None,ge=0,le=100),max_score:float|None=Query(default=None,ge=0,le=100),date_from:date|None=None,date_to:date|None=None):
    filters=dict(department=department,register_number=register_number,search=search,job_id=job_id,status=status,min_score=min_score,max_score=max_score,date_from=date_from,date_to=date_to)
    records=admin_application_query(db,**filters).options(joinedload(Application.student),joinedload(Application.job),joinedload(Application.result)).order_by(Application.created_at.asc()).all()
    screened=[a for a in records if a.result]; scores=[a.result.overall_score for a in screened]
    qualified=sum(a.result.qualification_status=='QUALIFIED' for a in screened); active_jobs=db.query(Job).filter(Job.active.is_(True))
    if job_id is not None: active_jobs=active_jobs.filter(Job.id==job_id)
    ranges=[(0,20),(21,40),(41,60),(61,80),(81,100)]
    distribution=[{'range':f'{low}–{high}','count':sum(low<=a.result.overall_score<=high for a in screened)} for low,high in ranges]
    departments={}
    for app_row in screened:
        label=app_row.student.department or 'Not provided'; data=departments.setdefault(label,{'department':label,'candidates':set(),'scores':[]}); data['candidates'].add(app_row.student_id); data['scores'].append(app_row.result.overall_score)
    department_analysis=[{'department':name,'candidates':len(data['candidates']),'screened':len(data['scores']),'average_score':round(sum(data['scores'])/len(data['scores']),1)} for name,data in departments.items()]
    jobs={}
    for app_row in screened:
        data=jobs.setdefault(app_row.job_id,{'job_id':app_row.job_id,'title':app_row.job.title,'screened':0,'scores':[]}); data['screened']+=1; data['scores'].append(app_row.result.overall_score)
    component_values={key:[] for key in ['required','optional','semantic_relevance','experience','projects','certifications','resume_quality','technical_expression']}
    trend=[]
    for app_row in screened:
        details=json.loads(app_row.result.details)
        for key,values in component_values.items():
            value=details.get('scores',{}).get(key)
            if isinstance(value,(int,float)): values.append(value)
        trend.append({'date':app_row.result.created_at.isoformat(),'score':app_row.result.overall_score,'student_id':app_row.student_id,'job_id':app_row.job_id})
    components=[{'component':key,'average_score':round(sum(values)/len(values),1),'count':len(values)} for key,values in component_values.items() if values]
    by_job=[{'job_id':v['job_id'],'title':v['title'],'screened':v['screened'],'count':v['screened'],'average_score':round(sum(v['scores'])/len(v['scores']),1)} for v in jobs.values()]
    application_restricted=any([job_id is not None,status,min_score is not None,max_score is not None,date_from,date_to,search])
    if application_restricted:
        candidate_count=len({a.student_id for a in records})
    else:
        users=db.query(User).filter(User.role=='student')
        if department: users=users.filter(func.lower(User.department)==department.lower())
        if register_number: users=users.filter(func.lower(User.register_number).contains(register_number.lower()))
        if search:
            term=f'%{search.strip().lower()}%'; users=users.filter(or_(func.lower(User.name).like(term),func.lower(User.email).like(term),func.lower(func.coalesce(User.register_number,'')).like(term),func.lower(func.coalesce(User.department,'')).like(term)))
        candidate_count=users.count()
    return {'candidates':candidate_count,'total_resumes_screened':len(screened),'qualified':qualified,'not_qualified':len(screened)-qualified,'average_score':round(sum(scores)/len(scores),1) if scores else None,'highest_score':max(scores) if scores else None,'active_job_roles':active_jobs.count(),'score_distribution':distribution,'qualification':{'qualified':qualified,'not_qualified':len(screened)-qualified},'department_analysis':department_analysis,'by_job':by_job,'component_performance':components,'trend':trend,'applications':len(records),'students':candidate_count,'jobs':active_jobs.count(),'screened':len(screened),'rejected':len(screened)-qualified,'filter_options':{'departments':[x[0] for x in db.query(User.department).filter(User.role=='student',User.department.is_not(None)).distinct().order_by(User.department).all()],'jobs':[{'id':j.id,'title':j.title} for j in db.query(Job).order_by(Job.title).all()]}}
def report_data(aid,db,u):
    a=db.get(Application,aid); owned(a,u)
    if not a.result: raise HTTPException(404,'Analysis is not available yet.')
    return a,json.loads(a.result.details)
@app.get('/api/reports/{aid}/docx')
def docx_report(aid:int,db:Session=Depends(get_db),u:User=Depends(current_user)):
    a,d=report_data(aid,db,u)
    try:
        from docx import Document
        doc=Document(); doc.add_heading('HIREX',0); doc.add_heading('ATS & Resume Screening Report',1)
        doc.add_paragraph(f'Candidate: {a.student.name} | Job: {a.job.title}'); doc.add_heading(f"Overall ATS Score: {d['scores']['required']*0 + a.result.overall_score:.0f}/100 — {d['qualification_status']}",2)
        for title,key in [('Score Breakdown','scores'),('Required Skill Analysis','required_skills'),('Optional Skill Analysis','optional_skills'),('Strengths','strengths'),('Weaknesses','weaknesses'),('Missing Skills','missing_skills'),('Recommendations','recommendations')]:
            doc.add_heading(title,2); obj=d.get(key,{})
            if isinstance(obj,dict): doc.add_paragraph(json.dumps(obj,indent=2))
            else:
                for item in obj: doc.add_paragraph(str(item),style='List Bullet')
        doc.add_heading('Qualification and Score Threshold',2)
        doc.add_paragraph(f"Minimum required: {d.get('minimum_required_score', 'Not configured')}/100 | Difference: {a.result.overall_score - int(d.get('minimum_required_score', 0)):+d} points | Status: {d.get('qualification_status', 'Unknown')}")
        for title,key in [('Education Eligibility','education_eligibility'),('Resume Quality','resume_quality_analysis'),('Technical Expression','technical_expression')]:
            doc.add_heading(title,2); doc.add_paragraph(json.dumps(d.get(key,{}),indent=2))
        doc.add_heading('Required Skill Evidence Matrix',2)
        for item in d.get('required_skills',[]):
            doc.add_paragraph(f"{item.get('skill')}: Exact evidence: {'Yes' if item.get('exact_match') else 'No'} | Demonstrated: {item.get('match_type') or 'No evidence'} | Related capabilities: {', '.join(item.get('related_capabilities',[])) or 'None'} | Evidence score: {item.get('evidence_score',0)}/100",style='List Bullet')
            if item.get('detailed_evidence'): doc.add_paragraph('Resume evidence: '+'; '.join(item['detailed_evidence']))
        doc.add_heading('Resume Section Evidence',2)
        for item in d.get('resume_section_evidence',[]):
            doc.add_paragraph(f"{item.get('section')}: {item.get('evidence_count',0)} evidence statements",style='List Bullet')
            for evidence in item.get('evidence',[]): doc.add_paragraph(evidence)
        doc.add_heading('Additional Skills Discovered',2)
        for item in d.get('additional_skills_found',[]): doc.add_paragraph(str(item),style='List Bullet')
        doc.add_heading('Project-Derived Knowledge',2)
        doc.add_paragraph(json.dumps(d.get('project_derived_knowledge',{}),indent=2))
        doc.add_heading('Project-by-Project Analysis',2)
        for project in d.get('project_analysis',{}).get('projects',[]):
            doc.add_heading(project.get('project_name','Project'),3)
            for label,key in [('Domain','project_domain'),('Languages','programming_languages'),('Frameworks','frameworks'),('Libraries','libraries'),('Tools','tools'),('Models and algorithms','models_and_algorithms'),('Methods','methods'),('Demonstrated capabilities','technical_concepts'),('Application area','application_area'),('Resume evidence','implementation_evidence'),('Reported results','results_metrics')]:
                value=project.get(key)
                if value: doc.add_paragraph(f"{label}: {value if isinstance(value,str) else ', '.join(value)}")
        doc.add_heading('Technical Evidence',2)
        for item in d.get('technical_evidence',[]):
            doc.add_paragraph(f"{item.get('source_section','Resume')} — {item.get('project_name') or 'Experience'}: {item.get('implementation','')}")
            if item.get('results_metrics'): doc.add_paragraph('Reported results: '+'; '.join(item['results_metrics']))
        stream=io.BytesIO(); doc.save(stream); stream.seek(0); return StreamingResponse(stream,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename=hirex-report-{aid}.docx'})
    except ImportError: raise HTTPException(501,'DOCX reporting dependency is not installed.')
@app.get('/api/reports/{aid}/pdf')
def pdf_report(aid:int,db:Session=Depends(get_db),u:User=Depends(current_user)):
    a,d=report_data(aid,db,u)
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        from reportlab.lib import colors
        stream=io.BytesIO(); c=canvas.Canvas(stream,pagesize=letter); width,height=letter; y=height-54
        def line(txt,size=10,bold=False):
            nonlocal y
            c.setFont('Helvetica-Bold' if bold else 'Helvetica',size)
            for wrapped in textwrap.wrap(str(txt),width=96) or ['']:
                if y<55: c.showPage(); y=height-54
                c.drawString(48,y,wrapped); y-=16
        c.setFillColor(colors.HexColor('#123a44')); line('HIREX',22,True); line('ATS & Resume Screening Report',14,True); line(f'Candidate: {a.student.name}  |  Job: {a.job.title}'); line(f"Overall ATS Score: {a.result.overall_score:.0f}/100 — {d['qualification_status']}",13,True)
        for title,key in [('Score Breakdown','scores'),('Required Skill Analysis','required_skills'),('Optional Skill Analysis','optional_skills'),('Strengths','strengths'),('Weaknesses','weaknesses'),('Missing Skills','missing_skills'),('Recommendations','recommendations')]:
            y-=6; line(title,12,True); obj=d.get(key,{})
            vals=[f'{k}: {v}' for k,v in obj.items()] if isinstance(obj,dict) else [str(x) for x in obj]
            for v in vals: line('• '+v)
        line('Qualification and Score Threshold',12,True)
        line(f"Minimum required: {d.get('minimum_required_score', 'Not configured')}/100 | Difference: {a.result.overall_score - int(d.get('minimum_required_score', 0)):+d} points | Status: {d.get('qualification_status', 'Unknown')}")
        for title,key in [('Education Eligibility','education_eligibility'),('Resume Quality','resume_quality_analysis'),('Technical Expression','technical_expression')]:
            y-=5; line(title,12,True)
            value=d.get(key,{})
            if isinstance(value,dict):
                for field,detail in value.items(): line(f'{field.replace("_"," ")}: {detail}')
            else: line(str(value))
        y-=5; line('Required Skill Evidence Matrix',12,True)
        for item in d.get('required_skills',[]):
            line(f"{item.get('skill')}: Exact: {'Yes' if item.get('exact_match') else 'No'} | Demonstrated: {item.get('match_type') or 'No evidence'} | Related: {', '.join(item.get('related_capabilities',[])) or 'None'} | Score: {item.get('evidence_score',0)}/100")
            for evidence in item.get('detailed_evidence',[]): line('  Evidence: '+evidence)
        y-=5; line('Resume Section Evidence',12,True)
        for item in d.get('resume_section_evidence',[]):
            line(f"{item.get('section')}: {item.get('evidence_count',0)} evidence statements")
            for evidence in item.get('evidence',[]): line('  '+evidence)
        y-=6; line('Additional Skills Discovered',12,True)
        for item in d.get('additional_skills_found',[]): line('• '+item)
        y-=6; line('Project-Derived Knowledge',12,True)
        for key,value in d.get('project_derived_knowledge',{}).items(): line(f'{key.replace("_"," ")}: {", ".join(value)}')
        y-=6; line('Project-by-Project Analysis',12,True)
        for project in d.get('project_analysis',{}).get('projects',[]):
            line(f"{project.get('project_name','Project')} — Evidence {project.get('evidence_score',0)}/100",10,True)
            for label,key in [('Domain','project_domain'),('Languages','programming_languages'),('Frameworks','frameworks'),('Libraries','libraries'),('Tools','tools'),('Models and algorithms','models_and_algorithms'),('Methods','methods'),('Demonstrated capabilities','technical_concepts'),('Application area','application_area'),('Resume evidence','implementation_evidence'),('Reported results','results_metrics')]:
                value=project.get(key)
                if value: line(f"{label}: {value if isinstance(value,str) else '; '.join(value)}")
        y-=6; line('Technical Evidence',12,True)
        for item in d.get('technical_evidence',[]):
            line(f"{item.get('source_section','Resume')} — {item.get('project_name') or 'Experience'}: {item.get('implementation','')}")
        c.save(); stream.seek(0); return StreamingResponse(stream,media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename=hirex-report-{aid}.pdf'})
    except ImportError: raise HTTPException(501,'PDF reporting dependency is not installed.')
